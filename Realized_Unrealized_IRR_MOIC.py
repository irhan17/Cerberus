#Imports libraries
import pandas as pd
import numpy as np
from scipy.optimize import brentq
from openpyxl import load_workbook
from openpyxl.styles import Font

#Loads the positions file
pos = pd.read_excel('2026-05-01_CMBS_Positions.xlsx',
                    sheet_name='Global_Asset_Detail', header=1, usecols=['CUSIP'])
pos.columns = pos.columns.str.strip()
pos_cusips = set(pos['CUSIP'].dropna().astype(str).str.strip().unique())
print(f"Loaded {len(pos_cusips)} CUSIPs from positions file")

#Loads the cash flow file
cf = pd.read_excel('CMBS_CF_MBS_ONLY.xlsx')
cf.columns = cf.columns.str.strip()
cf['Settle Date'] = pd.to_datetime(cf['Settle Date'])
cf['vCUSIP'] = cf['vCUSIP'].astype(str).str.strip()
cf['Derived Fund Family Group Short'] = (cf['Derived Fund Family Group Short'].astype(str).str.strip())
print(f"Loaded {len(cf)} rows from CF file")

#Defines the XIRR function
def xirr(dates, amounts):
    if len(dates) < 2:
        return None
    if all(a >= 0 for a in amounts) or all(a <= 0 for a in amounts):
        return None
    dates_num = np.array([(d - dates[0]).days / 365.25 for d in dates])
    amounts = np.array(amounts)
    def npv(rate):
        return np.sum(amounts / (1 + rate) ** dates_num)
    try:
        return brentq(npv, -0.5, 10.0, maxiter=1000)
    except:
        return None

#Defines the MOIC function
def moic(amounts):
    invested = sum(a for a in amounts if a < 0)
    returned = sum(a for a in amounts if a > 0)
    if invested == 0:
        return None
    return returned / abs(invested)

#Computes IRR and MOIC per (Fund, vCUSIP)
xirr_map = {}
moic_map = {}

for (fund, vcusip), group in cf.groupby(['Derived Fund Family Group Short', 'vCUSIP']):
    group = group.sort_values('Settle Date')
    dates = group['Settle Date'].dt.to_pydatetime().tolist()
    amounts = group['CashFlowAmount'].tolist()
    key = (fund, vcusip)
    xirr_map[key] = xirr(dates, amounts)
    moic_map[key] = moic(amounts)
print(f"Computed IRR/MOIC for {len(xirr_map)} Fund-vCUSIP pairs")

#Writes Realized vs Unrealized IRR / MOIC into Excel
wb = load_workbook('CMBS_CF_MBS_ONLY.xlsx')
ws = wb['MBS_CF']
ws['P1'] = 'Realized IRR'
ws['Q1'] = 'Realized MOIC'
ws['R1'] = 'Unrealized IRR'
ws['S1'] = 'Unrealized MOIC'
for col in ['P1', 'Q1', 'R1', 'S1']:
    ws[col].font = Font(bold=True)

#Populate row-level values
for row in range(2, ws.max_row + 1):
    fund = ws.cell(row=row, column=3).value   
    vcusip = ws.cell(row=row, column=5).value 
    if fund is None or vcusip is None:
        continue
    fund = str(fund).strip()
    vcusip = str(vcusip).strip()
    key = (fund, vcusip)
    irr = xirr_map.get(key)
    m = moic_map.get(key)

    #Unrealized
    if vcusip in pos_cusips:
        ws.cell(row=row, column=18, value=irr if irr is not None else "NA")
        ws.cell(row=row, column=19, value=m if m is not None else "NA")
        ws.cell(row=row, column=16, value="NA")
        ws.cell(row=row, column=17, value="NA")

    #Realized
    else:
        ws.cell(row=row, column=16, value=irr if irr is not None else "NA")
        ws.cell(row=row, column=17, value=m if m is not None else "NA")
        ws.cell(row=row, column=18, value="NA")
        ws.cell(row=row, column=19, value="NA")

#Saves file
wb.save('CMBS_CF_MBS_ONLY_with_Fund_Realized_Unrealized_IRR_MOIC.xlsx')
print("Program done - saved to CMBS_CF_MBS_ONLY_with_Fund_Realized_Unrealized_IRR_MOIC.xlsx")