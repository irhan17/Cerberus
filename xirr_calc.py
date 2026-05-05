#Imports libraries
import pandas as pd
import numpy as np
from scipy.optimize import brentq
from openpyxl import load_workbook
from openpyxl.styles import Font

#Loads the positions CUSIPs
pos = pd.read_excel('2026-05-01_CMBS_Positions.xlsx',
                     sheet_name='Global_Asset_Detail', header=1, usecols=['CUSIP'])
pos_cusips = set(pos['CUSIP'].dropna().unique())
print(f"Loaded {len(pos_cusips)} CUSIPs from positions file")

#Loads the cash flow file
cf = pd.read_excel('CMBS_CF_MBS_ONLY.xlsx',
                    usecols=['Settle Date', 'Cusip', 'CashFlowAmount'])
cf['Settle Date'] = pd.to_datetime(cf['Settle Date'])
print(f"Loaded {len(cf)} rows from CF file")

#Defines the XIRR function
def xirr(dates, amounts):
    if len(dates) < 2:
        return None
    if all(a >= 0 for a in amounts) or all(a <= 0 for a in amounts):
        return None
    dates_num = np.array([(d - dates[0]).days / 365.25 for d in dates])
    amounts = np.array(amounts)
    def npv(rate):
        return np.sum(amounts / (1 + rate) ** dates_num)
    try:
        return brentq(npv, -0.5, 10.0, maxiter=1000)
    except:
        return None

#Computes XIRR per CUSIP
cf_matched = cf[cf['Cusip'].isin(pos_cusips)]
print(f"Matched {cf_matched['Cusip'].nunique()} CUSIPs, {len(cf_matched)} rows")

xirr_map = {}
for cusip, group in cf_matched.groupby('Cusip'):
    group = group.sort_values('Settle Date')
    dates = group['Settle Date'].dt.to_pydatetime().tolist()
    amounts = group['CashFlowAmount'].tolist()
    xirr_map[cusip] = xirr(dates, amounts)

print(f"Computed XIRR for {sum(v is not None for v in xirr_map.values())} CUSIPs")

#Writes XIRR into column P of the CF file ---
wb = load_workbook('CMBS_CF_MBS_ONLY.xlsx')
ws = wb['MBS_CF']
ws['P1'] = 'XIRR'
ws['P1'].font = Font(bold=True)

for row in range(2, ws.max_row + 1):
    cusip = ws.cell(row=row, column=5).value 
    irr = xirr_map.get(cusip)
    if irr is not None:
        cell = ws.cell(row=row, column=16, value=irr)
        cell.number_format = '0.00%'

#Saves to file
wb.save('CMBS_CF_MBS_ONLY_with_XIRR.xlsx')
print("Program done - saved to CMBS_CF_MBS_ONLY_with_XIRR.xlsx")