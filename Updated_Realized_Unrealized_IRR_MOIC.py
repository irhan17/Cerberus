#Imports libraries
import pandas as pd
import numpy as np
from scipy.optimize import brentq
import shutil
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

#Loads the positions file
pos = pd.read_excel('2026-05-01_CMBS_Positions.xlsx',
                    sheet_name='Global_Asset_Detail', header=1, usecols=['CUSIP'])
pos.columns = pos.columns.str.strip()
pos_cusips = set(pos['CUSIP'].dropna().astype(str).str.strip().unique())
print(f"Loaded {len(pos_cusips)} CUSIPs from positions file")

#Loads the cash flow file
cf = pd.read_excel('CMBS_CF_MBS_ONLY.xlsx')
cf.columns = cf.columns.str.strip()
cf['Settle Date'] = pd.to_datetime(cf['Settle Date'])
cf['vCUSIP'] = cf['vCUSIP'].astype(str).str.strip()
cf['Derived Fund Family Group Short'] = cf['Derived Fund Family Group Short'].astype(str).str.strip()
print(f"Loaded {len(cf)} rows from CF file")

#Defines XIRR function
def xirr(dates, amounts):
    if len(dates) < 2:
        return None
    if all(a >= 0 for a in amounts) or all(a <= 0 for a in amounts):
        return None
    dates_num = np.array([(d - dates[0]).days / 365.25 for d in dates])
    amounts = np.array(amounts)
    def npv(rate):
        return np.sum(amounts / (1 + rate) ** dates_num)
    try:
        return brentq(npv, -0.5, 10.0, maxiter=1000)
    except:
        return None

#Defines MOIC function: Sum(All CF except Cost) / (-Cost)
def moic(df_group):
    cost = df_group.loc[df_group['Transaction Type Group'] == 'Cost', 'CashFlowAmount'].sum()
    other_cf = df_group.loc[df_group['Transaction Type Group'] != 'Cost', 'CashFlowAmount'].sum()
    if cost == 0:
        return None
    return other_cf / -cost

#Computes IRR and MOIC per (Fund, vCUSIP)
xirr_map = {}
moic_map = {}
for (fund, vcusip), group in cf.groupby(['Derived Fund Family Group Short', 'vCUSIP']):
    group = group.sort_values('Settle Date')
    dates = group['Settle Date'].dt.to_pydatetime().tolist()
    amounts = group['CashFlowAmount'].tolist()
    key = (fund, vcusip)
    xirr_map[key] = xirr(dates, amounts)
    moic_map[key] = moic(group)
print(f"Computed IRR/MOIC for {len(xirr_map)} Fund-vCUSIP pairs")

#Builds Return tab
rows = []
for (fund, vcusip), group in cf.groupby(['Derived Fund Family Group Short', 'vCUSIP']):
    key = (fund, vcusip)
    irr = xirr_map.get(key)
    m = moic_map.get(key)
    is_unrealized = vcusip in pos_cusips
    if is_unrealized:
        rows.append({'Derived Fund Family Group Short': fund,
                    'vCUSIP': vcusip,
                    'Security': group['Security'].iloc[0],
                    'Realized IRR': 'N/A',
                    'Realized MOIC': 'N/A',
                    'Unrealized IRR': irr if irr is not None else None,
                    'Unrealized MOIC': m})
    else:
        rows.append({'Derived Fund Family Group Short': fund,
                    'vCUSIP': vcusip,
                    'Security': group['Security'].iloc[0],
                    'Realized IRR': irr if irr is not None else None,
                    'Realized MOIC': m,
                    'Unrealized IRR': 'N/A',
                    'Unrealized MOIC': 'N/A'})
return_df = pd.DataFrame(rows)

#Copies original file and writes to Return worksheet
output_file = 'Updated_CMBS_CF_MBS_ONLY_with_Fund_Realized_Unrealized_IRR_MOIC.xlsx'
shutil.copy('CMBS_CF_MBS_ONLY.xlsx', output_file)
with pd.ExcelWriter(output_file, engine='openpyxl', mode='a', if_sheet_exists='replace') as writer:
    return_df.to_excel(writer, sheet_name='Return', index=False)

#Keeps same formatting from before
wb = load_workbook(output_file)
ws = wb['Return']
header_fill = PatternFill(start_color='4F81BD', end_color='4F81BD', fill_type='solid')
header_font = Font(bold=True, color='FFFFFF')
for cell in ws[1]:
    cell.fill = header_fill
    cell.font = header_font

#Saves file
wb.save(output_file)
print("Program done - saved to Updated_CMBS_CF_MBS_ONLY_with_Fund_Realized_Unrealized_IRR_MOIC.xlsx")